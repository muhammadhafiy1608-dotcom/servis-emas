from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = 'Senarai'
headers = ['Nama', 'No Telefon', 'Tarikh Lahir', 'No IC', 'Kumpulan Umur', 'Panggilan', 'Jantina', 'Kumpulan', 'Catatan']
widths = [34, 18, 16, 18, 18, 18, 12, 22, 34]
gold = PatternFill('solid', fgColor='F6ECD6')
plum = '43173A'
thin = Side(style='thin', color='D3C4CD')
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color=plum, size=12)
    c.fill = gold
    c.alignment = Alignment(vertical='center')
    c.border = Border(bottom=thin)
    ws.column_dimensions[c.column_letter].width = widths[i-1]
ws.row_dimensions[1].height = 24
ws.freeze_panes = 'A2'
# Telefon, Tarikh Lahir & IC sebagai teks supaya Excel tak buang 0 di depan / tukar format
for col in ('B', 'C', 'D'):
    for r in range(2, 2001):
        ws[f'{col}{r}'].number_format = '@'
dv = DataValidation(type='list', formula1='"Lelaki,Perempuan"', allow_blank=True)
dv.error = 'Pilih Lelaki atau Perempuan (atau biarkan kosong)'
ws.add_data_validation(dv)
dv.add('G2:G2000')

a = wb.create_sheet('Arahan')
a.column_dimensions['A'].width = 22
a.column_dimensions['B'].width = 70
rows = [
    ('Templat senarai penyimpan — Servis Emas', None),
    (None, None),
    ('Cara guna', 'Isi helaian "Senarai" (satu baris = satu penyimpan), simpan, kemudian import dalam app: Penyimpan → Import.'),
    (None, None),
    ('Lajur', 'Penerangan'),
    ('Nama *', 'Nama penuh seperti dalam IC. Contoh: Siti Aminah binti Ali'),
    ('No Telefon *', 'Nombor WhatsApp. Contoh: 012-345 6789 atau +60123456789'),
    ('Tarikh Lahir', 'Hari/bulan/tahun. Contoh: 16/08/1990. Tahun diperlukan untuk mesej ikut umur.'),
    ('No IC', 'Pilihan. Jika tarikh lahir kosong, app ambil tarikh lahir & jantina daripada IC. IC TIDAK disimpan dalam app.'),
    ('Kumpulan Umur', 'Pilihan, jika tiada tarikh lahir atau IC. Tulis umur (35), julat (30-39 atau 30an) atau nama kumpulan: Muda, Dewasa, Pertengahan, Pra-bersara, Warga emas. Menentukan set ilmu yang diterima.'),
    ('Panggilan', 'Pilihan. Contoh: Kak Ros, Abang Mat. Kosongkan untuk automatik (bawah 30: Hafiz; 30 ke atas: Tuan Ali / Puan Aminah).'),
    ('Jantina', 'Pilihan. Lelaki atau Perempuan. App boleh teka daripada bin/binti atau IC.'),
    ('Kumpulan', 'Pilihan. Pisahkan dengan koma. Contoh: VIP, Johor. Letak "Bukan Islam" untuk penyimpan bukan Islam.'),
    ('Catatan', 'Pilihan. Nota untuk rujukan anda sendiri.'),
    (None, None),
    ('Contoh baris', None),
]
for r in rows:
    a.append(list(r))
a['A1'].font = Font(bold=True, size=14, color=plum)
for r in (3, 5, 16):
    a.cell(row=r, column=1).font = Font(bold=True, color=plum)
a.cell(row=5, column=2).font = Font(bold=True, color=plum)
ex_head = headers
ex = [
    ['Siti Aminah binti Ali', '012-345 6789', '16/08/1990', '', '', '', '', 'VIP', 'Simpan untuk haji'],
    ['Muhammad Hafiz bin Rahman', '019-876 5432', '', '980312-10-5671', '', '', '', '', ''],
    ['Tan Mei Ling', '016-222 3344', '03/03/1983', '', '', 'Puan Tan', 'Perempuan', 'Bukan Islam', ''],
    ['Hasnah binti Omar', '013-456 7788', '', '', 'Warga emas', '', '', '', 'Tiada tarikh lahir, guna kumpulan'],
]
start = 17
for j, h in enumerate(ex_head, start=1):
    c = a.cell(row=start, column=j, value=h); c.font = Font(bold=True, color=plum); c.fill = gold
for i, row in enumerate(ex, start=start+1):
    for j, v in enumerate(row, start=1):
        a.cell(row=i, column=j, value=v)
for col, w in zip('CDEFGHI', [16, 18, 16, 14, 12, 16, 26]):
    a.column_dimensions[col].width = max(a.column_dimensions[col].width or 0, w)
for r in range(1, 30):
    a.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
wb.active = 0
import os
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'public', 'templat-senarai-penyimpan.xlsx'))
print('saved')
