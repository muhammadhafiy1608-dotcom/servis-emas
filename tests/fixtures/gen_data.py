import random, datetime, csv
from openpyxl import Workbook
random.seed(7)
male = ['Ahmad','Muhammad','Mohd','Hafiz','Faiz','Amir','Hakim','Rizal','Zulkifli','Rosli','Azman','Hisham','Firdaus','Syafiq','Aiman','Haziq','Kamarul','Ismail','Rahman','Nizam','Fauzi','Shahrul','Idris','Yusof','Daud']
female = ['Siti','Nurul','Aminah','Aisyah','Farah','Hasnah','Rosnah','Noraini','Haslinda','Azlina','Faridah','Zalina','Syafiqah','Izzah','Hidayah','Najwa','Liyana','Amira','Salmah','Rohani']
fathers = ['Ali','Abu','Hassan','Ismail','Ahmad','Abdullah','Omar','Yusof','Rahman','Salleh','Kamal','Hamzah','Zainal','Daud','Mokhtar','Rashid']
cn = ['Tan Mei Ling','Lim Chee Keong','Lee Wei Ming','Wong Siew Lan','Ng Kok Leong','Chong Ai Ling','Ong Boon Hock','Teh Swee Lin','Goh Chin Yee','Chan Kah Mun']
ind = [('Kumar','L','Raju'),('Kavitha','P','Subramaniam'),('Suresh','L','Muthu'),('Devi','P','Krishnan'),('Rajesh','L','Gopal')]
today = datetime.date(2026,10,28)
rows=[]
for i in range(600):
    r = random.random()
    tags=''
    if r < 0.08:
        name = random.choice(cn); g = None; tags='Bukan Islam'
    elif r < 0.12:
        f, g, fa = random.choice(ind); name = f'{f} {"a/l" if g=="L" else "a/p"} {fa}'; tags='Bukan Islam'
    elif r < 0.56:
        g='L'; pre=random.choice(['', 'Muhammad ', 'Mohd ', 'Ahmad ', 'Wan ', '']); name=f'{pre}{random.choice(male)} bin {random.choice(fathers)}'
    else:
        g='P'; pre=random.choice(['', 'Siti ', 'Nurul ', 'Nor ', '']); name=f'{pre}{random.choice(female)} binti {random.choice(fathers)}'
    if random.random() < 0.25: name = name.upper()
    age = random.choice(range(19, 78))
    dob = datetime.date(today.year-age, random.randint(1,12), random.randint(1,28))
    if i in (3, 17, 250): dob = datetime.date(today.year-age, 10, 28)  # hari jadi hari ini
    yy = dob.strftime('%y%m%d')
    last = random.randint(0,9999)
    if g=='L' and last%2==0: last+=1
    if g=='P' and last%2==1: last+=1
    last%=10000
    ic = f'{yy}-{random.choice(["01","08","10","14","03"])}-{last:04d}'
    mode = random.random()
    dob_s=''; ic_s=''
    if mode < 0.55: dob_s = dob.strftime('%d/%m/%Y')
    elif mode < 0.80: ic_s = ic
    elif mode < 0.88: dob_s = dob.strftime('%-d %b %Y')
    # else: tiada tarikh lahir
    n = 10000000 + i*37
    prefix = random.choice(['012','013','017','019','011'])
    if prefix=='011': local = f'011{n % 100000000:08d}'
    else: local = f'{prefix}{n % 10000000:07d}'
    fm = random.random()
    if fm < 0.4: ph = f'{local[:3]}-{local[3:6]} {local[6:]}'
    elif fm < 0.6: ph = local
    elif fm < 0.75: ph = '+6' + local
    elif fm < 0.85: ph = '6' + local
    else: ph = local[1:]  # Excel buang 0
    if i in (100, 101): ph = ''
    if i == 102: ph = 'tiada'
    rows.append([name, ph, dob_s, ic_s, '', '', tags, ''])
# pendua: baris 200 ulang nombor baris 199 dgn nama lain (anak)
rows[200][1] = rows[199][1]; rows[200][0] = 'Adam bin Ali'
wb=Workbook(); ws=wb.active; ws.title='Pelanggan'
ws.append(['Bil','Nama Penuh','No. Tel','Tarikh Lahir','No IC','Panggilan','Jantina','Kumpulan','Catatan','No Akaun GAP'])
for i,r in enumerate(rows, start=1):
    ws.append([i] + r + [f'GAP{100000+i}'])
# sebahagian nombor telefon sebagai nombor (bukan teks) — seperti Excel sebenar
for rr in range(2, 40):
    v = ws.cell(row=rr, column=3).value
    if v and v.isdigit(): ws.cell(row=rr, column=3).value = int(v)
# sebahagian tarikh lahir sebagai tarikh Excel sebenar
for rr in range(40, 80):
    v = ws.cell(row=rr, column=4).value
    if v and '/' in v:
        d,m,y = map(int, v.split('/')); c = ws.cell(row=rr, column=4); c.value = datetime.date(y,m,d); c.number_format='dd/mm/yyyy'
import os
base = os.path.dirname(os.path.abspath(__file__)) + '/'
wb.save(base+'penyimpan-600.xlsx')
with open(base+'penyimpan-600.csv','w',newline='',encoding='utf-8') as f:
    w=csv.writer(f); w.writerow(['Nama','No Telefon','Tarikh Lahir','No IC','Panggilan','Jantina','Kumpulan','Catatan']); w.writerows(rows)
print('rows', len(rows))
